"""Read results from the Unics 100plusV over SSH, read-only.

The analyzer saves each result as a row of /home/pi/Desktop/31-10-2020/result.xlsx
(A test, B method, C result, D unit, E date, F patient ID, G blank OD,
H standard OD, I sample OD, J patient name). Ranges come from the same
analyzer's test setup in Patientdata.xlsx (Test Low / Test High). Both are
opened read-only; nothing on the analyzer is written. The SSH password lives
in the macOS Keychain or the Windows Credential Manager, never in a file.

On macOS the system ssh is used (it keeps one connection open between
reads). On Windows a built-in Python SSH client (paramiko) is used instead,
so the lab PC needs no SSH setup. VETLAB_SSH=paramiko forces it anywhere.
"""

import json
import os
import re
import shlex
import socket
import subprocess
import sys
import threading
from pathlib import Path

PROGRAM_DIR = "/home/pi/Desktop/31-10-2020"
KEYCHAIN_SERVICE = "VetLab Unics100plusV"

REMOTE_SCRIPT = r"""
import json, os, sys, time
import openpyxl
program_dir, known = sys.argv[1], sys.argv[2]
def read(name):
    try:
        return open(os.path.join(program_dir, name)).read().strip()
    except OSError:
        return ''
results = os.path.join(program_dir, 'result.xlsx')
setup = os.path.join(program_dir, 'Patientdata.xlsx')
mtime = '{}|{}'.format(os.stat(results).st_mtime, os.stat(setup).st_mtime)
out = {'mtime': mtime, 'next_row': read('resultstore.txt'), 'next_patient_id': read('ptid.txt'),
       'rows': None, 'ranges': None}
if mtime != known:
    for attempt in range(3):
        try:
            sheet = openpyxl.load_workbook(results, read_only=True, data_only=True).active
            names = ('test_name', 'method', 'result', 'unit', 'date', 'patient_id',
                     'blank_od', 'std_od', 'sample_od', 'patient_name')
            rows = []
            for number, values in enumerate(sheet.iter_rows(min_row=2, max_col=10, values_only=True), 2):
                if values and values[0] is not None:
                    row = {'row': number}
                    row.update({k: '' if v is None else str(v) for k, v in zip(names, values)})
                    rows.append(row)
            book = openpyxl.load_workbook(setup, read_only=True, data_only=True)['Pdata']
            table = list(book.iter_rows(values_only=True))
            head = [str(h or '').strip().lower() for h in table[0]]
            col = lambda name: head.index(name) if name in head else None
            pick = {'test_low': col('test low'), 'test_high': col('test high'), 'unit': col('unit'),
                    'linearity_low': col('linearity low'), 'linearity_high': col('linearity high'),
                    'method': col('test method')}
            ranges = {}
            for values in table[1:]:
                name = values[col('test name')] if col('test name') is not None else None
                if name:
                    ranges[str(name)] = {k: ('' if i is None or values[i] is None else str(values[i]))
                                         for k, i in pick.items()}
            out.update(rows=rows, ranges=ranges)
            break
        except Exception as exc:
            out['error'] = str(exc)
            time.sleep(1)
print(json.dumps(out))
"""


class AnalyzerError(RuntimeError):
    pass


# ---------------------------------------------------------------- password

WINDOWS = sys.platform == "win32"


def get_password(user):
    if WINDOWS:
        import keyring
        return keyring.get_password(KEYCHAIN_SERVICE, user) or ""
    found = subprocess.run(
        ["security", "find-generic-password", "-s", KEYCHAIN_SERVICE, "-a", user, "-w"],
        capture_output=True, text=True,
    )
    return found.stdout.rstrip("\n") if found.returncode == 0 else ""


def set_password(user, password):
    if WINDOWS:
        import keyring
        keyring.set_password(KEYCHAIN_SERVICE, user, password)
        return
    subprocess.run(
        ["security", "add-generic-password", "-U", "-s", KEYCHAIN_SERVICE, "-a", user, "-w", password],
        capture_output=True, text=True, check=True,
    )


# ---------------------------------------------------------------- ssh

def _askpass(folder):
    path = Path(folder) / "askpass.sh"
    if not path.exists():
        path.write_text('#!/bin/sh\nprintf "%s\\n" "$VETLAB_SSH_PASSWORD"\n')
        path.chmod(0o700)
    return path


def _hidden():
    """Keep console windows from flashing up on Windows."""
    return {"creationflags": subprocess.CREATE_NO_WINDOW} if WINDOWS else {}


def machine_id(host):
    """MAC address of whichever unit answers at host (all units are raspberrypi.local)."""
    try:
        ip = socket.gethostbyname(host)
    except OSError:
        return ""
    ping = ["ping", "-n", "1", "-w", "2000", ip] if WINDOWS else ["ping", "-c", "1", "-t", "2", ip]
    subprocess.run(ping, capture_output=True, **_hidden())
    arp = subprocess.run(["arp", "-a" if WINDOWS else "-n", ip], capture_output=True, text=True, **_hidden()).stdout
    # macOS: "? (169.254.1.2) at 2c:cf:67:6:e3:f4 on en8"; Windows: "  169.254.1.2   2c-cf-67-06-e3-f4   dynamic"
    found = re.search(r"([0-9a-f]{1,2}(?:[:-][0-9a-f]{1,2}){5})", arp.lower())
    return "".join(part.zfill(2) for part in re.split(r"[:-]", found.group(1))) if found else ""


def fetch(host, user, password, work_folder, known_mtime=""):
    """Return {'mtime', 'next_row', 'next_patient_id', 'rows', 'ranges', 'unit_mac'}; rows is None if unchanged."""
    if not password:
        raise AnalyzerError("No 100plusV password saved. Enter it in Instruments settings.")
    unit = machine_id(host) or "unknown"
    if WINDOWS or os.environ.get("VETLAB_SSH") == "paramiko":
        data = _fetch_paramiko(host, user, password, work_folder, known_mtime, unit)
        data["unit_mac"] = unit
        return data
    options = [
        "-o", "ConnectTimeout=5",
        "-o", "HostKeyAlias=unics-{}".format(unit),
        "-o", "StrictHostKeyChecking=accept-new",
        "-o", "ControlMaster=auto",
        "-o", "ControlPath={}/.ssh/vetlab-{}-%r:%p".format(Path.home(), unit),
        "-o", "ControlPersist=120",
    ]
    env = dict(os.environ, VETLAB_SSH_PASSWORD=password, SSH_ASKPASS=str(_askpass(work_folder)),
               SSH_ASKPASS_REQUIRE="force", DISPLAY=os.environ.get("DISPLAY", ":0"))
    target = "{}@{}".format(user, host)
    for attempt in range(2):
        try:
            done = subprocess.run(
                # ssh hands the command to the analyzer's shell as one line, so quote the
                # arguments (the mtime marker contains "|").
                ["ssh", *options, target, "python3", "-", shlex.quote(PROGRAM_DIR), shlex.quote(known_mtime or "-")],
                input=REMOTE_SCRIPT, text=True, capture_output=True, timeout=30, env=env,
            )
        except subprocess.TimeoutExpired:
            raise AnalyzerError("No answer from {} within 30 s".format(host))
        if done.returncode == 0:
            break
        lines = done.stderr.strip().splitlines()
        reason = lines[-1] if lines else "ssh to {} failed".format(host)
        if "Permission denied" in reason:
            raise AnalyzerError(reason + " (wrong password)")
        if "Connection refused" in reason:
            raise AnalyzerError(reason + " (SSH is off on the analyzer)")
        if "resolve" in reason:
            raise AnalyzerError(reason + " (analyzer not on the network)")
        # A dropped connection (e.g. BrokenPipeError on the analyzer side):
        # close the shared connection and try once more on a fresh one.
        subprocess.run(["ssh", *options, "-O", "exit", target], capture_output=True, env=env)
    else:
        if "BrokenPipe" in reason or "closed" in reason or "reset" in reason:
            reason = "Connection to the analyzer dropped twice in a row. Check the cable; retrying in 30 s."
        raise AnalyzerError(reason)
    data = json.loads(done.stdout)
    data["unit_mac"] = unit
    return data


# ---------------------------------------------------------------- ssh without the system ssh (Windows)

_clients = {}
_clients_lock = threading.Lock()


def _client(host, user, password, work_folder, unit):
    """One open connection per unit. Like ssh's accept-new, each unit's key is remembered
    the first time (by its hardware address) and a changed key is refused."""
    import paramiko

    with _clients_lock:
        client = _clients.get(unit)
        if client is not None and client.get_transport() is not None and client.get_transport().is_active():
            return client
        client = paramiko.SSHClient()
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())  # checked below against our own file
        try:
            client.connect(host, username=user, password=password, timeout=5, banner_timeout=10,
                           auth_timeout=10, look_for_keys=False, allow_agent=False)
        except paramiko.AuthenticationException:
            raise AnalyzerError("Permission denied (wrong password)")
        except (socket.gaierror, socket.herror):
            raise AnalyzerError("Cannot find {} (analyzer not on the network)".format(host))
        except ConnectionRefusedError:
            raise AnalyzerError("Connection refused (SSH is off on the analyzer)")
        except (OSError, paramiko.SSHException) as exc:
            raise AnalyzerError("Cannot reach {}: {}".format(host, exc))
        key = client.get_transport().get_remote_server_key()
        known_file = Path(work_folder) / "known_hosts"
        known = paramiko.HostKeys(str(known_file)) if known_file.exists() else paramiko.HostKeys()
        alias = "unics-{}".format(unit)
        stored = known.lookup(alias)
        if stored and key.get_name() in stored and stored[key.get_name()] != key:
            client.close()
            raise AnalyzerError("This analyzer's SSH key changed. If the analyzer was replaced, delete "
                                "{} and try again.".format(known_file))
        if not stored or key.get_name() not in stored:
            known.add(alias, key.get_name(), key)
            known.save(str(known_file))
        client.get_transport().set_keepalive(15)
        _clients[unit] = client
        return client


def _fetch_paramiko(host, user, password, work_folder, known_mtime, unit):
    import paramiko

    command = "python3 - {} {}".format(shlex.quote(PROGRAM_DIR), shlex.quote(known_mtime or "-"))
    reason = ""
    for attempt in range(2):
        client = _client(host, user, password, work_folder, unit)
        try:
            stdin, stdout, stderr = client.exec_command(command, timeout=30)
            stdin.write(REMOTE_SCRIPT)
            stdin.channel.shutdown_write()
            out = stdout.read().decode()
            err = stderr.read().decode()
            if stdout.channel.recv_exit_status() == 0:
                return json.loads(out)
            lines = err.strip().splitlines()
            reason = lines[-1] if lines else "the read on the analyzer failed"
        except (OSError, EOFError, paramiko.SSHException) as exc:
            reason = str(exc) or type(exc).__name__
        # Drop the connection and try once more on a fresh one.
        with _clients_lock:
            _clients.pop(unit, None)
        client.close()
    if "BrokenPipe" in reason or "closed" in reason or "reset" in reason or "timed out" in reason:
        reason = "Connection to the analyzer dropped twice in a row. Check the cable; retrying in 30 s."
    raise AnalyzerError(reason)
